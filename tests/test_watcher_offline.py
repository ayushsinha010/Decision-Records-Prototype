"""Offline tests for watcher.py.

These never call the Anthropic API: a stub client returns a canned JSON
reply, so the transcript reading, JSON parsing, Markdown rendering and
processed-file tracking can be checked without a key or network access.

Run from the project root:

    python -m unittest discover -s tests -v
"""

import json
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import watcher  # noqa: E402
from docx import Document  # noqa: E402

SAMPLE = Path(__file__).resolve().parent.parent / "samples" / "sample_meeting.txt"

CANNED_FIELDS = {
    "date_or_meeting": "Weekly Pipeline Review - 3 March 2026",
    "category": "Account Planning",
    "question_considered": "How should near-certain deals be classified?",
    "assumptions": "A verbal commitment is a reasonable proxy for a signed PO.",
    "evidence_reviewed": "",
    "decision_made": "- Reclassify near-certain deals with a verbal commitment as confident.",
    "resulting_action": "- Update the pipeline sheet by Friday.",
    "owner": "Jordan Blake",
    "status": "Open",
    "account_client": "Client A plc",
    "stakeholders_mentioned": "Alex Morgan\nJordan Blake\nMaya Fernandes",
    "requirements_mentioned": "",
    "products_discussed": "Reporting dashboard",
    "open_questions": "Can billing data be exported weekly?",
}


class StubClient:
    """Mimics the one call watcher.py makes: client.messages.create(...)."""

    def __init__(self, reply_text):
        self.reply_text = reply_text
        self.calls = []
        self.messages = SimpleNamespace(create=self._create)

    def _create(self, **kwargs):
        self.calls.append(kwargs)
        block = SimpleNamespace(type="text", text=self.reply_text)
        return SimpleNamespace(content=[block])


class ReadTranscriptTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def test_txt_is_read_and_stripped(self):
        p = self.tmp / "a.txt"
        p.write_text("\n  hello world  \n", encoding="utf-8")
        self.assertEqual(watcher.read_transcript_text(p), "hello world")

    def test_vtt_header_and_cues_are_removed(self):
        p = self.tmp / "a.vtt"
        p.write_text(
            "WEBVTT\n\n1\n00:00:01.000 --> 00:00:03.500\nAlex Morgan: Hello.\n\n"
            "2\n00:00:04.000 --> 00:00:06.000\nJordan Blake: Hi.\n",
            encoding="utf-8",
        )
        text = watcher.read_transcript_text(p)
        self.assertNotIn("WEBVTT", text)
        self.assertNotIn("-->", text)
        self.assertIn("Alex Morgan: Hello.", text)
        self.assertIn("Jordan Blake: Hi.", text)

    def test_docx_paragraphs_are_joined_and_blank_ones_dropped(self):
        p = self.tmp / "a.docx"
        doc = Document()
        doc.add_paragraph("First line")
        doc.add_paragraph("   ")
        doc.add_paragraph("Second line")
        doc.save(str(p))
        self.assertEqual(watcher.read_transcript_text(p), "First line\nSecond line")

    def test_extension_check_is_case_insensitive(self):
        p = self.tmp / "A.TXT"
        p.write_text("loud file", encoding="utf-8")
        self.assertEqual(watcher.read_transcript_text(p), "loud file")


class ExtractAndRenderTests(unittest.TestCase):
    def test_plain_json_reply(self):
        client = StubClient(json.dumps(CANNED_FIELDS))
        self.assertEqual(watcher.extract_fields(client, "text"), CANNED_FIELDS)
        call = client.calls[0]
        self.assertEqual(call["model"], watcher.MODEL)
        self.assertEqual(call["system"], watcher.SYSTEM_PROMPT)
        self.assertIn("Transcript:", call["messages"][0]["content"])

    def test_fenced_json_reply_is_unwrapped(self):
        for fence in ("```json", "```"):
            client = StubClient(f"{fence}\n{json.dumps(CANNED_FIELDS)}\n```")
            self.assertEqual(watcher.extract_fields(client, "text"), CANNED_FIELDS)

    def test_non_json_reply_raises(self):
        client = StubClient("Sorry, I cannot do that.")
        with self.assertRaises(json.JSONDecodeError):
            watcher.extract_fields(client, "text")

    def test_render_markdown_contains_every_section(self):
        md = watcher.render_markdown(CANNED_FIELDS, "sample_meeting.txt")
        self.assertTrue(md.startswith("# Weekly Pipeline Review - 3 March 2026"))
        for needle in (
            "**Source transcript:** `sample_meeting.txt`",
            "## Decision",
            "## Context",
            "**Category:** Account Planning",
            "**Owner:** Jordan Blake",
            "**Account/Client:** Client A plc",
            "Can billing data be exported weekly?",
        ):
            self.assertIn(needle, md)

    def test_render_markdown_empty_fields_show_dash_and_filename_title(self):
        md = watcher.render_markdown({}, "empty.txt")
        self.assertTrue(md.startswith("# empty.txt"))
        self.assertIn("**Decision Made:**\n-\n", md)
        self.assertIn("**Open Questions:**\n-\n", md)


class ProcessFileTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)
        self.transcripts = self.tmp / "transcripts_in"
        self.output = self.tmp / "librechat_file_search"
        self.transcripts.mkdir()
        self.log_path = self.tmp / ".processed.json"
        self._patches = [
            mock.patch.object(watcher, "TRANSCRIPTS_FOLDER", self.transcripts),
            mock.patch.object(watcher, "OUTPUT_FOLDER", self.output),
            mock.patch.object(watcher, "PROCESSED_LOG", self.log_path),
            mock.patch.object(watcher, "wait_until_stable", lambda *_a, **_k: True),
        ]
        for p in self._patches:
            p.start()

    def tearDown(self):
        for p in self._patches:
            p.stop()
        self._tmp.cleanup()

    def _drop_sample(self, name="sample_meeting.txt"):
        dest = self.transcripts / name
        dest.write_text(SAMPLE.read_text(encoding="utf-8"), encoding="utf-8")
        return dest

    def test_sample_transcript_end_to_end_with_stub_client(self):
        self.assertTrue(SAMPLE.exists(), "sample transcript is missing")
        path = self._drop_sample()
        client = StubClient(json.dumps(CANNED_FIELDS))
        processed = set()

        watcher.process_file(path, client, processed)

        out = self.output / "sample_meeting.md"
        self.assertTrue(out.exists())
        self.assertIn("**Category:** Account Planning", out.read_text(encoding="utf-8"))
        self.assertEqual(len(client.calls), 1)
        self.assertIn("Client A plc renewal", client.calls[0]["messages"][0]["content"])
        self.assertEqual(json.loads(self.log_path.read_text(encoding="utf-8")), ["sample_meeting.txt"])
        self.assertEqual(watcher.load_processed(), {"sample_meeting.txt"})

        # A second pass over the same file must not call the model again.
        watcher.process_file(path, client, processed)
        self.assertEqual(len(client.calls), 1)

    def test_unsupported_extension_is_ignored(self):
        path = self.transcripts / "notes.pdf"
        path.write_bytes(b"%PDF-1.4")
        client = StubClient(json.dumps(CANNED_FIELDS))
        watcher.process_file(path, client, set())
        self.assertEqual(client.calls, [])
        self.assertFalse(self.output.exists())

    def test_empty_transcript_is_skipped(self):
        path = self.transcripts / "empty.txt"
        path.write_text("   \n", encoding="utf-8")
        client = StubClient(json.dumps(CANNED_FIELDS))
        processed = set()
        watcher.process_file(path, client, processed)
        self.assertEqual(client.calls, [])
        self.assertEqual(processed, set())

    def test_invalid_json_is_not_marked_processed(self):
        path = self._drop_sample()
        client = StubClient("not json at all")
        processed = set()
        with self.assertLogs("decision-records-watcher", level="ERROR"):
            watcher.process_file(path, client, processed)
        self.assertEqual(processed, set())
        self.assertFalse(self.log_path.exists())
        self.assertFalse((self.output / "sample_meeting.md").exists())


class WaitUntilStableTests(unittest.TestCase):
    def test_missing_file_returns_false(self):
        with tempfile.TemporaryDirectory() as d:
            self.assertFalse(watcher.wait_until_stable(Path(d) / "nope.txt", check_seconds=0))

    def test_existing_file_returns_true(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "a.txt"
            p.write_text("x", encoding="utf-8")
            self.assertTrue(watcher.wait_until_stable(p, check_seconds=0))


if __name__ == "__main__":
    unittest.main()
